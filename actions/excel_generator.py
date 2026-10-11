"""
actions/excel_generator.py — Generates an Excel spreadsheet (.xlsx) and opens it.
Runs in the background; announces completion on its own. Uses
core/text_generation.py, which tries Gemini first and falls back to a local
LLM (Ollama) if Gemini is unreachable, with tolerant JSON parsing for smaller
local models.
"""
import importlib.util
import json
import re
import subprocess
import sys
import threading
import time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def _base_dir() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).parent
    return Path(__file__).resolve().parent.parent


def _config() -> dict:
    try:
        return json.loads((_base_dir() / "config" / "api_keys.json").read_text(encoding="utf-8"))
    except Exception:
        return {}


def _get_os() -> str:
    return _config().get("os_system", "windows").lower()


def _output_dir() -> Path:
    out = _base_dir() / "generated_sheets"
    out.mkdir(parents=True, exist_ok=True)
    return out


def _open_file(path: Path) -> None:
    os_name = _get_os()
    try:
        if os_name == "windows":
            import os
            os.startfile(str(path))
        elif os_name == "mac":
            subprocess.run(["open", str(path)], check=False)
        else:
            subprocess.run(["xdg-open", str(path)], check=False)
    except Exception as e:
        print(f"[ExcelGenerator] ⚠️ Could not auto-open {path}: {e}")


def _say(player, text: str) -> None:
    if player and hasattr(player, "request_say"):
        try:
            player.request_say(text)
        except Exception:
            pass


def _get_text_generation_module():
    key = "jarvis_ext_text_generation"
    if key in sys.modules:
        return sys.modules[key]
    path = _base_dir() / "core" / "text_generation.py"
    spec = importlib.util.spec_from_file_location(key, str(path))
    mod = importlib.util.module_from_spec(spec)
    sys.modules[key] = mod
    spec.loader.exec_module(mod)
    return mod


_SHEET_SYSTEM = (
    "You output ONLY valid JSON, nothing else — no markdown fences, no commentary. "
    'Format: {"title": "<sheet title>", "headers": ["Col1", "Col2", ...], '
    '"rows": [["val1", "val2", ...], ...]}. Include 5-10 example/starter rows.'
)


def _get_sheet_data(description: str) -> dict:
    tg = _get_text_generation_module()
    prompt = f"{_SHEET_SYSTEM}\n\nCreate a spreadsheet for: {description}."
    data = tg.generate_json(prompt)
    if "headers" not in data or "rows" not in data:
        raise ValueError("Model returned an unexpected sheet shape.")
    return data


def _build_xlsx(data: dict) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = (data.get("title") or "Sheet1")[:31]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1C8CA8", end_color="1C8CA8", fill_type="solid")

    for col_idx, header in enumerate(data["headers"], start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill

    for row_idx, row in enumerate(data["rows"], start=2):
        for col_idx, val in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=val)

    for col_idx in range(1, len(data["headers"]) + 1):
        ws.column_dimensions[chr(64 + col_idx)].width = 18

    safe_name = re.sub(r"[^\w\-]+", "_", data.get("title", "sheet")).strip("_")[:50] or "sheet"
    path = _output_dir() / f"{safe_name}_{int(time.time())}.xlsx"
    wb.save(str(path))
    return path


def _work(description: str, player, session_memory):
    try:
        data = _get_sheet_data(description)
        path = _build_xlsx(data)
    except Exception as e:
        err = f"Sir, I couldn't build that spreadsheet: {e}"
        _log(err, player)
        _say(player, err)
        return

    _open_file(path)
    done = "Your spreadsheet is ready, sir, and it's open now."
    _log(done, player)
    _say(player, done)
    if session_memory:
        try:
            session_memory.set_last_search(query=f"generate spreadsheet: {description}", response=done)
        except Exception:
            pass


def generate_excel_sheet(parameters: dict, player=None, session_memory=None) -> str:
    description = (parameters or {}).get("description", "").strip()
    if not description:
        msg = "Sir, what should the spreadsheet track?"
        _log(msg, player)
        return msg

    threading.Thread(target=_work, args=(description, player, session_memory), daemon=True).start()

    msg = "On it, sir — building that spreadsheet now. I'll let you know when it's ready."
    _log(msg, player)
    return msg


def _log(message: str, player=None) -> None:
    print(f"[ExcelGenerator] {message}")
    if player:
        try:
            player.write_log(f"JARVIS: {message}")
        except Exception:
            pass


TOOL = {
    "name": "generate_excel_sheet",
    "description": (
        "Creates an Excel spreadsheet in the background and opens it when ready. "
        "Returns immediately with an acknowledgment — the file arrives as a "
        "follow-up, spoken on its own."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "description": {"type": "STRING", "description": "What the spreadsheet should track, e.g. 'monthly budget with categories'"}
        },
        "required": ["description"]
    },
    "handler": generate_excel_sheet,
}