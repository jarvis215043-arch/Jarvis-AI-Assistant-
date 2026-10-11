"""
actions/word_generator.py — Generates a Word document (.docx) on a topic and
opens it. Runs in the background; announces completion on its own. Uses
core/text_generation.py, which tries Gemini first and falls back to a local
LLM (Ollama) if Gemini is unreachable, with tolerant JSON parsing for smaller
local models.
"""
import importlib.util
import json
import re
import subprocess
import sys
import threading
import time
from pathlib import Path

from docx import Document
from docx.shared import Pt


def _base_dir() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).parent
    return Path(__file__).resolve().parent.parent


def _config() -> dict:
    try:
        return json.loads((_base_dir() / "config" / "api_keys.json").read_text(encoding="utf-8"))
    except Exception:
        return {}


def _get_os() -> str:
    return _config().get("os_system", "windows").lower()


def _output_dir() -> Path:
    out = _base_dir() / "generated_docs"
    out.mkdir(parents=True, exist_ok=True)
    return out


def _open_file(path: Path) -> None:
    os_name = _get_os()
    try:
        if os_name == "windows":
            import os
            os.startfile(str(path))
        elif os_name == "mac":
            subprocess.run(["open", str(path)], check=False)
        else:
            subprocess.run(["xdg-open", str(path)], check=False)
    except Exception as e:
        print(f"[WordGenerator] ⚠️ Could not auto-open {path}: {e}")


def _say(player, text: str) -> None:
    if player and hasattr(player, "request_say"):
        try:
            player.request_say(text)
        except Exception:
            pass


def _get_text_generation_module():
    key = "jarvis_ext_text_generation"
    if key in sys.modules:
        return sys.modules[key]
    path = _base_dir() / "core" / "text_generation.py"
    spec = importlib.util.spec_from_file_location(key, str(path))
    mod = importlib.util.module_from_spec(spec)
    sys.modules[key] = mod
    spec.loader.exec_module(mod)
    return mod


_DOC_SYSTEM = (
    "You output ONLY valid JSON, nothing else — no markdown fences, no commentary. "
    'Format: {"title": "<document title>", "sections": [{"heading": "<section heading>", '
    '"body": "<paragraph text, 3-6 sentences>"}, ...]}'
)


def _get_content(topic: str, num_sections: int) -> dict:
    tg = _get_text_generation_module()
    prompt = (
        f"{_DOC_SYSTEM}\n\n"
        f"Write a {num_sections}-section document on: {topic}. "
        f"Each section needs a heading and a solid paragraph of body text."
    )
    data = tg.generate_json(prompt)
    if "title" not in data or "sections" not in data:
        raise ValueError("Model returned an unexpected document shape.")
    return data


def _build_docx(content: dict) -> Path:
    doc = Document()
    doc.add_heading(content["title"], level=0)
    for section in content["sections"]:
        doc.add_heading(section.get("heading", ""), level=1)
        p = doc.add_paragraph(section.get("body", ""))
        p.style.font.size = Pt(11)

    safe_name = re.sub(r"[^\w\-]+", "_", content["title"]).strip("_")[:50] or "document"
    path = _output_dir() / f"{safe_name}_{int(time.time())}.docx"
    doc.save(str(path))
    return path


def _work(topic: str, num_sections: int, player, session_memory):
    try:
        content = _get_content(topic, num_sections)
        path    = _build_docx(content)
    except Exception as e:
        err = f"Sir, I couldn't build that document: {e}"
        _log(err, player)
        _say(player, err)
        return

    _open_file(path)
    done = f"Your document on {topic} is ready, sir, and it's open now."
    _log(done, player)
    _say(player, done)
    if session_memory:
        try:
            session_memory.set_last_search(query=f"generate document: {topic}", response=done)
        except Exception:
            pass


def generate_word_doc(parameters: dict, player=None, session_memory=None) -> str:
    params       = parameters or {}
    topic        = (params.get("topic") or "").strip()
    num_sections = max(2, min(int(params.get("num_sections") or 4), 10))

    if not topic:
        msg = "Sir, what should the document be about?"
        _log(msg, player)
        return msg

    threading.Thread(target=_work, args=(topic, num_sections, player, session_memory), daemon=True).start()

    msg = f"On it, sir — writing that document on {topic} now. I'll let you know when it's ready."
    _log(msg, player)
    return msg


def _log(message: str, player=None) -> None:
    print(f"[WordGenerator] {message}")
    if player:
        try:
            player.write_log(f"JARVIS: {message}")
        except Exception:
            pass


TOOL = {
    "name": "generate_word_doc",
    "description": (
        "Creates a Word document on a given topic in the background and opens it "
        "when ready. Returns immediately with an acknowledgment — the file "
        "arrives as a follow-up, spoken on its own."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "topic": {"type": "STRING", "description": "Topic or title for the document"},
            "num_sections": {"type": "NUMBER", "description": "Number of sections (default 4)"}
        },
        "required": ["topic"]
    },
    "handler": generate_word_doc,
}